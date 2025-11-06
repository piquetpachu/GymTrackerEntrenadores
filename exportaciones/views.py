from progreso.models import Progreso, ProgresoEjercicio
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from clientes.models import Cliente
from django.http import HttpResponse

bold_font = Font(bold=True, size=12, color="FFFFFF")
title_font = Font(bold=True, size=14)
header_fill = PatternFill("solid", fgColor="4F81BD")
center_align = Alignment(horizontal="center", vertical="center")
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

def exportar_progreso_excel(request, cliente_id):
    cliente = Cliente.objects.get(pk=cliente_id)
    progresos = Progreso.objects.filter(cliente=cliente).order_by("fecha")
    progresos_ejercicios = ProgresoEjercicio.objects.filter(cliente=cliente).select_related("ejercicio", "rutina")

    wb = Workbook()

    # --- HOJA 1: Progreso físico ---
    ws1 = wb.active
    ws1.title = "Progreso físico"

    ws1.merge_cells("A1:I1")
    ws1["A1"] = f"Progreso físico de {cliente.nombre}"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = center_align

    headers = ["Fecha", "Peso (kg)", "Altura (cm)", "IMC", "Cintura", "Pecho", "Brazo", "Pierna", "Notas"]
    ws1.append(headers)

    for col in range(1, len(headers) + 1):
        cell = ws1.cell(row=2, column=col)
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for p in progresos:
        fila = [
            p.fecha.strftime("%d-%m-%Y"),
            p.peso_kg,
            p.altura_cm,
            round(p.peso_kg / ((p.altura_cm / 100) ** 2), 2) if p.altura_cm and p.peso_kg else "",
            p.cintura_cm,
            p.pecho_cm,
            p.brazo_cm,
            p.pierna_cm,
            p.notas,
        ]
        ws1.append(fila)
        for col in range(1, len(fila) + 1):
            cell = ws1.cell(row=ws1.max_row, column=col)
            cell.alignment = center_align
            cell.border = thin_border

    for i, col in enumerate(ws1.columns, 1):
        max_length = 0
        column = get_column_letter(i)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws1.column_dimensions[column].width = max_length + 2

    # --- HOJA 2: Progreso de rendimiento ---
    ws2 = wb.create_sheet(title="Rendimiento")

    ws2.merge_cells("A1:H1")
    ws2["A1"] = f"Rendimiento de {cliente.nombre}"
    ws2["A1"].font = title_font
    ws2["A1"].alignment = center_align

    headers2 = ["Fecha", "Rutina", "Ejercicio", "Series", "Reps", "Peso (kg)", "RPE", "RIR", "Notas"]
    ws2.append(headers2)

    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=2, column=col)
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for e in progresos_ejercicios:
        fila = [
            e.fecha.strftime("%d-%m-%Y"),
            e.rutina.nombre if e.rutina else "-",
            e.ejercicio.nombre,
            e.series,
            e.repeticiones,
            e.peso,
            e.rpe if e.rpe else "",
            e.rir if e.rir else "",
            e.notas,
        ]
        ws2.append(fila)
        for col in range(1, len(fila) + 1):
            cell = ws2.cell(row=ws2.max_row, column=col)
            cell.alignment = center_align
            cell.border = thin_border

    for i, col in enumerate(ws2.columns, 1):
        max_length = 0
        column = get_column_letter(i)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws2.column_dimensions[column].width = max_length + 2

    # --- Exportar archivo ---
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="Progreso_Completo_{cliente.nombre}.xlsx"'
    wb.save(response)
    return response

def exportar_rutina_excel(request, rutina_id):
    rutina = Rutina.objects.get(pk=rutina_id)
    entradas = EntradaEjercicio.objects.filter(rutina=rutina).select_related("ejercicio")

    wb = Workbook()
    ws = wb.active
    ws.title = rutina.nombre[:30]

    # Título
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Rutina: {rutina.nombre}"
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align

    ws.append([])

    headers = ["Ejercicio", "Series", "Repeticiones", "Descanso (segundos)", "Notas"]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col)
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for entrada in entradas:
        fila = [
            entrada.ejercicio.nombre,
            entrada.series,
            entrada.repeticiones,
            entrada.descanso_segundos,
            entrada.notas,
        ]
        ws.append(fila)

        for col in range(1, len(fila) + 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.alignment = center_align
            cell.border = thin_border

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="Rutina_{rutina.nombre}.xlsx"'
    wb.save(response)
    return response
